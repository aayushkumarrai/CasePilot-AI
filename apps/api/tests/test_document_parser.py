from io import BytesIO

import pytest
from docx import Document

from app.modules.documents.parser import DocumentExtractionError, extract_passages


def test_txt_passages_are_stable_and_split_long_text() -> None:
    data = ("First block.\n\n" + ("word " * 500)).encode()
    passages = extract_passages(data, "text/plain")
    assert [passage.sequence_number for passage in passages] == list(range(1, len(passages) + 1))
    assert passages[0].passage_label == "Passage 1"
    assert all(len(passage.content) <= 1500 for passage in passages)


def test_docx_uses_non_empty_paragraphs() -> None:
    document = Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("   ")
    document.add_paragraph("Second paragraph")
    output = BytesIO()
    document.save(output)
    passages = extract_passages(output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert [passage.passage_label for passage in passages] == ["Paragraph 1", "Paragraph 2"]
    assert [passage.content for passage in passages] == ["First paragraph", "Second paragraph"]


def test_empty_or_non_utf8_text_fails_safely() -> None:
    with pytest.raises(DocumentExtractionError, match="No readable text"):
        extract_passages(b" \n\n ", "text/plain")
    with pytest.raises(DocumentExtractionError, match="UTF-8"):
        extract_passages(b"\xff\xfe", "text/plain")
